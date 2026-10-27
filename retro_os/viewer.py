"""Text extraction, web downloading, and terminal-friendly pagination for modern documents."""

from __future__ import annotations

import pathlib
import shutil
import ssl
import subprocess
import tempfile
import textwrap
import urllib.parse
import urllib.request
from html.parser import HTMLParser


class SimpleHTMLTextExtractor(HTMLParser):
    """Strip HTML markup and return clean, readable text formatted for retro terminals."""

    def __init__(self):
        super().__init__()
        self.text_parts = []
        self._ignore = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head", "noscript"):
            self._ignore = True
        elif tag in ("p", "br", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr"):
            self.text_parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head", "noscript"):
            self._ignore = False
        elif tag in ("p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "tr"):
            self.text_parts.append("\n")

    def handle_data(self, data):
        if not self._ignore:
            self.text_parts.append(data)

    def get_text(self):
        raw = "".join(self.text_parts)
        lines = [line.strip() for line in raw.splitlines()]
        cleaned = []
        for line in lines:
            if line:
                cleaned.append(line)
            elif cleaned and cleaned[-1] != "":
                cleaned.append("")
        return "\n".join(cleaned)


def fetch_url(url, dest_dir=None, dest_name=None, timeout=15):
    """Download a file from an HTTP/HTTPS URL and save it to disk.

    Returns (dest_path, byte_count).
    """
    url = str(url).strip()
    if "://" not in url:
        url = "https://" + url
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise ValueError(f"Unsupported URL scheme '{parsed.scheme}'; expected http or https.")

    req = urllib.request.Request(
        url,
        headers={"User-Agent": "RetroSerialHub/2.6 (Vintage Terminal Host)"},
    )
    ctx = ssl.create_default_context()
    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
        data = resp.read()
        ctype = resp.headers.get("Content-Type", "").split(";")[0].strip().lower()

    if not dest_name:
        url_path = parsed.path.rstrip("/")
        dest_name = pathlib.Path(url_path).name if url_path else ""
        if not dest_name or "." not in dest_name:
            ext_map = {
                "text/plain": ".txt",
                "text/html": ".html",
                "application/pdf": ".pdf",
                "application/json": ".json",
                "text/csv": ".csv",
                "text/markdown": ".md",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
            }
            ext = ext_map.get(ctype, ".txt")
            dest_name = (dest_name or "download") + ext

    target_folder = pathlib.Path(dest_dir) if dest_dir else pathlib.Path.cwd()
    target_folder.mkdir(parents=True, exist_ok=True)
    dest_path = target_folder / dest_name
    dest_path.write_bytes(data)
    return dest_path, len(data)


def extract_text_from_url(url, timeout=15):
    """Download a document or web page from an HTTP/HTTPS URL and extract readable text."""
    url = str(url).strip()
    if "://" not in url:
        url = "https://" + url
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise ValueError(f"Unsupported URL scheme '{parsed.scheme}'; expected http or https.")

    req = urllib.request.Request(
        url,
        headers={"User-Agent": "RetroSerialHub/2.6 (Vintage Terminal Host)"},
    )
    ctx = ssl.create_default_context()
    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
        data = resp.read()
        ctype = resp.headers.get("Content-Type", "").split(";")[0].strip().lower()

    url_path = parsed.path.rstrip("/")
    suffix = pathlib.Path(url_path).suffix.lower()
    if not suffix:
        ext_map = {
            "text/plain": ".txt",
            "text/html": ".html",
            "application/pdf": ".pdf",
            "application/json": ".json",
            "text/csv": ".csv",
            "text/markdown": ".md",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
        }
        suffix = ext_map.get(ctype, ".txt")

    if suffix in (".pdf", ".docx", ".doc", ".xlsx", ".xls"):
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(data)
            tmp_path = pathlib.Path(tmp.name)
        try:
            return extract_text(tmp_path)
        finally:
            try:
                tmp_path.unlink()
            except OSError:
                pass

    if suffix in (".html", ".htm") or "html" in ctype:
        raw_html = data.decode("utf-8", errors="replace")
        parser = SimpleHTMLTextExtractor()
        try:
            parser.feed(raw_html)
            text = parser.get_text()
            if text.strip():
                return text
        except Exception:
            pass
        return raw_html

    return data.decode("utf-8", errors="replace")


def extract_text(path):
    path = pathlib.Path(path)
    suffix = path.suffix.lower()
    if suffix in (".html", ".htm"):
        raw_html = path.read_text(encoding="utf-8", errors="replace")
        parser = SimpleHTMLTextExtractor()
        try:
            parser.feed(raw_html)
            text = parser.get_text()
            if text.strip():
                return text
        except Exception:
            pass
        return raw_html
    if suffix in (".txt", ".md", ".csv", ".tsv", ".log", ".json", ".xml"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF viewing requires the optional 'pypdf' package.") from exc
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX viewing requires the optional 'python-docx' package.") from exc
        document = Document(str(path))
        rows = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            rows.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(rows)
    if suffix == ".doc":
        antiword = shutil.which("antiword")
        if antiword:
            result = subprocess.run(
                [antiword, str(path)],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            return result.stdout
        raise RuntimeError("Legacy DOC viewing requires the 'antiword' utility; convert the file to DOCX.")
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX viewing requires the optional 'openpyxl' package.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            output = []
            for sheet in workbook.worksheets:
                output.append(f"[{sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    output.append(",".join("" if value is None else str(value) for value in row))
            return "\n".join(output)
        finally:
            workbook.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Legacy XLS viewing requires the optional 'xlrd' package.") from exc
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            output = []
            for sheet in workbook.sheets():
                output.append(f"[{sheet.name}]")
                output.extend(
                    ",".join(str(value) for value in sheet.row_values(row))
                    for row in range(sheet.nrows)
                )
            return "\n".join(output)
        finally:
            workbook.release_resources()
    raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")


def paginate_text(text, columns=80, page_lines=22):
    if columns not in (40, 80):
        raise ValueError("Terminal width must be 40 or 80 columns.")
    lines = []
    for source_line in text.splitlines() or [""]:
        if not source_line:
            lines.append("")
        else:
            lines.extend(textwrap.wrap(source_line, width=columns, replace_whitespace=False, drop_whitespace=False) or [""])
    return [lines[index:index + page_lines] for index in range(0, len(lines), page_lines)] or [[]]
